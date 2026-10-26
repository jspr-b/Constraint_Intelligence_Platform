"""Generate the mapping-wizard test exports (ADR-0011) from the synthetic plants' clean exports.

    uv run python scripts/gen_wizard_cases.py           # write
    uv run python scripts/gen_wizard_cases.py --check   # fail when a file is stale

Each case rewrites one plant's clean ERP export in the style of another ERP or MES: other file
and column names, languages, separators, date and number formats, codes, and layout. The plant
and its true mapping are known, so each case's `expected.yaml` (file roles, column matches, code
tables, formats, flow) is derived here mechanically from the plant's built-in mapping and its
standard event log, and its answer is the plant's known answer. Nothing is calculated by the
engine, and `expected.yaml` is never edited by hand to fit the wizard.
"""

import argparse
import csv
import io
import sys
from collections import Counter, defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from itertools import pairwise
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SYN = ROOT / "data" / "synthetic"
OUT = SYN / "wizard_cases"
MAPPINGS = ROOT / "packages" / "connectors" / "mappings"
RULE = "ADR-0011. Derived from the plant's true mapping before the wizard exists; never edited."
TOP_ROUTES = 6
XLSX_EPOCH = datetime(2026, 10, 26)  # fixed workbook timestamps, so --check is stable


@dataclass
class FileSpec:
    """How one source table of the plant's export is written in the case."""

    name: str
    """Output file name, or ``<workbook>.xlsx#<sheet>`` for a sheet in a workbook."""
    columns: dict[str, str] = field(default_factory=dict)
    """Source column → case column. Unlisted columns keep their name."""


@dataclass
class Case:
    name: str
    style: str
    plant: str
    mapping: str
    files: dict[str, FileSpec]
    """Source table → how it is written. Tables not listed are left out of the case."""
    delimiter: str = ","
    decimal: str = "."
    datetime_format: str = "%Y-%m-%d %H:%M"
    split_datetime: tuple[str, str] | None = None
    """(date format, time format): each datetime column becomes ``<col> date`` + ``<col> time``
    columns named by ``split_names``."""
    split_names: dict[str, tuple[str, str]] = field(default_factory=dict)
    recode: dict[str, dict[str, str]] = field(default_factory=dict)
    """Code table name → plant code → case code."""
    title_lines: tuple[str, ...] = ()
    footer: bool = False
    bom: bool = False
    extra_files: dict[str, str] = field(default_factory=dict)
    """File name → content of a file that is not used, with its reason in ``not_used``."""
    not_used: dict[str, str] = field(default_factory=dict)
    bookings: dict[str, str] | None = None
    """Hour bookings instead of the event log: column names for job, operation, step, start
    date, start time, duration and quantity."""
    answer: dict[str, Any] | str = "plant_answer"
    notes: tuple[str, ...] = ()


# --- the cases ----------------------------------------------------------------------------------

CASES: tuple[Case, ...] = (
    Case(
        name="exact_en",
        style="Exact Globe-like export in English: comma separated, ISO dates, decimal point",
        plant="plant_a",
        mapping="generic",
        delimiter=",",
        decimal=".",
        datetime_format="%Y-%m-%d %H:%M",
        recode={
            "event_status": {
                "GESTART": "STARTED",
                "GEREED": "FINISHED",
                "OMSTEL_START": "SETUP BEGIN",
                "OMSTEL_EINDE": "SETUP END",
                "STORING_START": "BREAKDOWN",
                "STORING_EINDE": "REPAIRED",
                "AFKEUR": "REJECTED",
                "AANKOMST": "ARRIVED",
                "VERTREK": "DEPARTED",
            },
            "order_status": {
                "Open": "Open",
                "Vrijgegeven": "Released",
                "In bewerking": "In process",
                "Verzonden": "Delivered",
            },
            "calendar_kind": {
                "Dienst": "Shift",
                "Pauze": "Break",
                "Onderhoud": "Maintenance",
                "Feestdag": "Public holiday",
            },
        },
        files={
            "werkplekken": FileSpec(
                "Resources.csv",
                {"Werkplek": "Resource", "Omschrijving": "Description", "Afdeling": "Department"},
            ),
            "artikelen": FileSpec(
                "Items.csv",
                {
                    "Artikel": "Item code",
                    "Omschrijving": "Description",
                    "Productgroep": "Item group",
                    "Kleurgroep": "Colour group",
                    "Verkoopprijs": "Sales price",
                    "Variabele kosten": "Variable cost",
                },
            ),
            "bewerkingen": FileSpec(
                "Routings.csv",
                {
                    "Artikel": "Item code",
                    "Bewerking nr": "Step no",
                    "Bewerking": "Operation",
                    "Werkplek": "Resource",
                    "Alternatief": "Alternative resource",
                    "Insteltijd (min)": "Setup time (min)",
                    "Stuktijd (min)": "Run time per unit (min)",
                },
            ),
            "orders": FileSpec(
                "ProductionOrders.csv",
                {
                    "Ordernummer": "Order number",
                    "Artikel": "Item code",
                    "Aantal": "Quantity",
                    "Leverdatum": "Delivery date",
                    "Status": "Status",
                    "Vrijgegeven": "Released on",
                    "Verzonden": "Delivered on",
                    "Materiaal gereed": "Material available",
                },
            ),
            "mes_events": FileSpec(
                "ShopFloorEvents.csv",
                {
                    "Werkplek": "Resource",
                    "Afdeling": "Department",
                    "Bon": "Order number",
                    "Bewerking": "Step no",
                    "Status": "Event",
                    "Tijdstip": "Timestamp",
                    "Aantal": "Good qty",
                    "Afkeur": "Rejected qty",
                    "Storingscode": "Breakdown code",
                    "Afkeurcode": "Reject code",
                    "Medewerker": "Employee",
                },
            ),
            "rooster": FileSpec(
                "Roster.csv",
                {
                    "Werkplek": "Resource",
                    "Type": "Type",
                    "Start": "From",
                    "Einde": "To",
                    "Overgenomen": "Covered",
                },
            ),
            "overuren": FileSpec(
                "Overtime.csv",
                {"Werkplek": "Resource", "Jaar": "Year", "Week": "Week", "Overuren": "Hours"},
            ),
            "capaciteitskosten": FileSpec(
                "ResourceCosts.csv",
                {
                    "Werkplek": "Resource",
                    "Overwerk per uur": "Overtime rate",
                    "Extra dienst": "Extra shift cost",
                },
            ),
            "omsteltijden": FileSpec(
                "Changeovers.csv",
                {
                    "Werkplek": "Resource",
                    "Van kleurgroep": "From colour",
                    "Naar kleurgroep": "To colour",
                    "Omsteltijd (min)": "Changeover (min)",
                },
            ),
            "uitbesteding": FileSpec(
                "Subcontracting.csv", {"Bewerking": "Operation", "Prijs per stuk": "Unit price"}
            ),
        },
        extra_files={
            "Customers.csv": "Customer code,Customer name,Country\n"
            "C001,Customer 001 (synthetic),NL\nC002,Customer 002 (synthetic),DE\n"
        },
        not_used={"Customers.csv": "Customer list: not needed to find what limits output."},
    ),
    Case(
        name="sap_de",
        style="SAP PP-like export in German: tab separated, date and time in two columns, "
        "decimal comma, own work centre codes",
        plant="plant_c",
        mapping="generic",
        delimiter="\t",
        decimal=",",
        split_datetime=("%d.%m.%Y", "%H:%M:%S"),
        split_names={
            "Tijdstip": ("Ist-Datum", "Ist-Uhrzeit"),
            "Start": ("Beginn-Datum", "Beginn-Uhrzeit"),
            "Einde": ("Ende-Datum", "Ende-Uhrzeit"),
            "Leverdatum": ("Liefertermin", "Lieferzeit"),
            "Vrijgegeven": ("Freigabe-Datum", "Freigabe-Uhrzeit"),
            "Verzonden": ("Versand-Datum", "Versand-Uhrzeit"),
            "Materiaal gereed": ("Material-Datum", "Material-Uhrzeit"),
        },
        recode={
            "resource": {
                "SNIJ1": "1000-LASER",
                "LAS1": "1000-SCHW1",
                "LAS2": "1000-SCHW2",
                "LAK1": "1000-LACK",
                "MONT1": "1000-MONT",
                "INPAK1": "1000-VERP",
            },
            "event_status": {
                "GESTART": "START",
                "GEREED": "ENDE",
                "OMSTEL_START": "RUEST_A",
                "OMSTEL_EINDE": "RUEST_E",
                "STORING_START": "STOER_A",
                "STORING_EINDE": "STOER_E",
                "AFKEUR": "AUSS",
                "AANKOMST": "ANK",
                "VERTREK": "ABG",
            },
            "order_status": {
                "Open": "ERÖF",
                "Vrijgegeven": "FREI",
                "In bewerking": "TRÜCK",
                "Verzonden": "GLFT",
            },
            "calendar_kind": {
                "Dienst": "Schicht",
                "Pauze": "Pause",
                "Onderhoud": "Wartung",
                "Feestdag": "Feiertag",
            },
        },
        files={
            "werkplekken": FileSpec(
                "CRHD_Arbeitsplaetze.txt",
                {
                    "Werkplek": "Arbeitsplatz",
                    "Omschrijving": "Kurztext",
                    "Afdeling": "Kostenstelle",
                },
            ),
            "artikelen": FileSpec(
                "MARA_Material.txt",
                {
                    "Artikel": "Material",
                    "Omschrijving": "Materialkurztext",
                    "Productgroep": "Warengruppe",
                    "Kleurgroep": "Farbgruppe",
                    "Verkoopprijs": "Verkaufspreis",
                    "Variabele kosten": "Variable Kosten",
                },
            ),
            "bewerkingen": FileSpec(
                "PLPO_Arbeitsplan.txt",
                {
                    "Artikel": "Material",
                    "Bewerking nr": "Vorgang",
                    "Bewerking": "Vorgangstext",
                    "Werkplek": "Arbeitsplatz",
                    "Alternatief": "Alternativ-Arbeitsplatz",
                    "Insteltijd (min)": "Ruestzeit (MIN)",
                    "Stuktijd (min)": "Bearbeitungszeit je Stk (MIN)",
                },
            ),
            "orders": FileSpec(
                "AFKO_Auftraege.txt",
                {
                    "Ordernummer": "Auftrag",
                    "Artikel": "Material",
                    "Aantal": "Auftragsmenge",
                    "Status": "Systemstatus",
                },
            ),
            "mes_events": FileSpec(
                "AFRU_Rueckmeldungen.txt",
                {
                    "Werkplek": "Arbeitsplatz",
                    "Afdeling": "Kostenstelle",
                    "Bon": "Auftrag",
                    "Bewerking": "Vorgang",
                    "Status": "Rueckmeldeart",
                    "Aantal": "Gutmenge",
                    "Afkeur": "Ausschuss",
                    "Storingscode": "Stoerursache",
                    "Afkeurcode": "Ausschussgrund",
                    "Medewerker": "Personalnummer",
                },
            ),
            "rooster": FileSpec(
                "KAKO_Kapazitaet.txt",
                {"Werkplek": "Arbeitsplatz", "Type": "Art", "Overgenomen": "Uebernommen"},
            ),
            "overuren": FileSpec(
                "Ueberstunden.txt",
                {"Werkplek": "Arbeitsplatz", "Jaar": "Jahr", "Week": "KW", "Overuren": "Stunden"},
            ),
            "capaciteitskosten": FileSpec(
                "Kostensaetze.txt",
                {
                    "Werkplek": "Arbeitsplatz",
                    "Overwerk per uur": "Ueberstundensatz",
                    "Extra dienst": "Zusatzschicht",
                },
            ),
            "omsteltijden": FileSpec(
                "Ruestmatrix.txt",
                {
                    "Werkplek": "Arbeitsplatz",
                    "Van kleurgroep": "Von Farbe",
                    "Naar kleurgroep": "Nach Farbe",
                    "Omsteltijd (min)": "Ruestzeit (MIN)",
                },
            ),
            "uitbesteding": FileSpec(
                "Fremdvergabe.txt", {"Bewerking": "Vorgangstext", "Prijs per stuk": "Preis je Stk"}
            ),
        },
    ),
    Case(
        name="isah_xlsx",
        style="Isah-like export: one Excel workbook with a sheet per table, real Excel dates "
        "and numbers, Dutch names",
        plant="plant_b",
        mapping="generic",
        files={
            "werkplekken": FileSpec(
                "isah_export.xlsx#Machines",
                {"Werkplek": "Machinecode", "Omschrijving": "Machine", "Afdeling": "Kostenplaats"},
            ),
            "artikelen": FileSpec(
                "isah_export.xlsx#Artikelen",
                {
                    "Artikel": "Artikelcode",
                    "Productgroep": "Artikelgroep",
                    "Kleurgroep": "Lakgroep",
                    "Verkoopprijs": "Verkoopprijs excl",
                    "Variabele kosten": "Kostprijs variabel",
                },
            ),
            "bewerkingen": FileSpec(
                "isah_export.xlsx#Bewerkingen",
                {
                    "Artikel": "Artikelcode",
                    "Bewerking nr": "Volgnr",
                    "Bewerking": "Omschrijving bewerking",
                    "Werkplek": "Machinecode",
                    "Alternatief": "Uitwijkmachine",
                    "Insteltijd (min)": "Insteltijd",
                    "Stuktijd (min)": "Bewerkingstijd per stuk",
                },
            ),
            "orders": FileSpec(
                "isah_export.xlsx#Productieorders",
                {
                    "Ordernummer": "Productieorder",
                    "Artikel": "Artikelcode",
                    "Aantal": "Order aantal",
                    "Leverdatum": "Gevraagde leverdatum",
                    "Vrijgegeven": "Vrijgavedatum",
                    "Verzonden": "Afleverdatum",
                    "Materiaal gereed": "Materiaal binnen",
                },
            ),
            "mes_events": FileSpec(
                "isah_export.xlsx#Terugmeldingen",
                {
                    "Werkplek": "Machinecode",
                    "Afdeling": "Kostenplaats",
                    "Bon": "Productieorder",
                    "Bewerking": "Volgnr",
                    "Status": "Meldingssoort",
                    "Tijdstip": "Datum tijd",
                    "Aantal": "Aantal goed",
                    "Afkeur": "Aantal afkeur",
                    "Storingscode": "Storingsoorzaak",
                    "Afkeurcode": "Afkeuroorzaak",
                    "Medewerker": "Personeelsnr",
                },
            ),
            "rooster": FileSpec(
                "isah_export.xlsx#Capaciteit",
                {"Werkplek": "Machinecode", "Type": "Soort", "Start": "Van", "Einde": "Tot"},
            ),
            "overuren": FileSpec("isah_export.xlsx#Overuren", {"Werkplek": "Machinecode"}),
            "capaciteitskosten": FileSpec("isah_export.xlsx#Tarieven", {"Werkplek": "Machinecode"}),
            "omsteltijden": FileSpec("isah_export.xlsx#Omstellen", {"Werkplek": "Machinecode"}),
            "uitbesteding": FileSpec("isah_export.xlsx#Uitbesteed"),
        },
        notes=("Dates and numbers are real Excel values, not text; order numbers are numbers.",),
    ),
    Case(
        name="mes_us",
        style="US-style MES export: month/day/year dates, English event words, a badge column",
        plant="plant_d",
        mapping="workshop_scans",
        datetime_format="%m/%d/%Y %H:%M",
        recode={
            "action": {
                "START": "Clock on",
                "FINISH": "Clock off",
                "STOP": "Interrupted",
                "RESTART": "Resumed",
                "ARRIVE": "Received",
                "LEAVE": "Dispatched",
            },
            "step_kind": {"Machine": "Equipment", "Handwork": "Manual"},
        },
        files={
            "workstations": FileSpec(
                "work_centers.csv",
                {
                    "Workstation": "Work Center",
                    "Description": "Work Center Name",
                    "Area code": "Dept",
                    "Area": "Dept Name",
                    "Type": "Category",
                    "Positions": "Stations",
                },
            ),
            "shift_plan": FileSpec(
                "schedule.csv",
                {
                    "Workstation": "Work Center",
                    "Type": "Entry",
                    "Start": "Begins",
                    "End": "Ends",
                    "Staffed": "Crew slots",
                },
            ),
            "teams": FileSpec(
                "crews.csv", {"Team": "Crew", "Name": "Crew Name", "Workstations": "Work Centers"}
            ),
            "team_plan": FileSpec(
                "crew_schedule.csv",
                {"Team": "Crew", "Start": "Begins", "End": "Ends", "Staffed": "Crew size"},
            ),
            "jobs": FileSpec(
                "work_orders.csv",
                {
                    "Job": "Work Order",
                    "Qty": "Quantity",
                    "Due": "Due Date",
                    "Status": "WO Status",
                    "Released": "Released",
                    "Shipped": "Completed",
                    "Order value": "Sales Value",
                    "Material cost": "Material Cost",
                },
            ),
            "scans": FileSpec(
                "labor_transactions.csv",
                {
                    "Job": "Work Order",
                    "Workstation": "Work Center",
                    "Area code": "Dept",
                    "Time": "Transaction Time",
                    "Action": "Transaction",
                    "Qty": "Quantity",
                    "Reason": "Reason Code",
                    "Operator": "Badge",
                },
            ),
        },
        notes=("Month/day order is only certain from days above 12 (e.g. 07/20/2026).",),
    ),
    Case(
        name="ridder_hours",
        style="Ridder iQ-like hour bookings only: start date, start time and duration per job and "
        "step, no scans, setups or stops",
        plant="plant_a",
        mapping="generic",
        delimiter=";",
        decimal=",",
        datetime_format="%d-%m-%Y %H:%M",
        files={
            "werkplekken": FileSpec(
                "Werkplekken.csv",
                {"Werkplek": "Werkplekcode", "Omschrijving": "Werkplek", "Afdeling": "Afdeling"},
            ),
            "artikelen": FileSpec("Artikelen.csv"),
            "bewerkingen": FileSpec("Bewerkingen.csv"),
            "orders": FileSpec("Productieorders.csv", {"Ordernummer": "Productieorder"}),
            "rooster": FileSpec("Rooster.csv", {"Werkplek": "Werkplekcode"}),
        },
        bookings={
            "file": "Urenregistratie.csv",
            "job": "Productieorder",
            "operation": "Bewerkingsnr",
            "step": "Werkplekcode",
            "date": "Datum",
            "time": "Begintijd",
            "duration": "Duur (min)",
            "quantity": "Gereed gemeld",
        },
        answer={
            "constraint": "PNT-01",
            "confidence": "lower than high, and said in plain words: times come from hour bookings",
            "not_compared": "lost euros and loss causes: bookings have no setups, stops or scans",
        },
    ),
    Case(
        name="messy",
        style="Messy workshop export: report title lines above the header, a totals line at the "
        "end, a byte-order mark, semicolons, and the plan tables together in one workbook",
        plant="plant_d",
        mapping="workshop_scans",
        delimiter=";",
        title_lines=("Werkplaatsrapport scans", "Geëxporteerd op 26-10-2026 08:15", ""),
        footer=True,
        bom=True,
        files={
            "workstations": FileSpec("rapport_werkplekken.csv"),
            "jobs": FileSpec("rapport_jobs.csv"),
            "scans": FileSpec("rapport_scans.csv"),
            "shift_plan": FileSpec("planning.xlsx#Rooster"),
            "teams": FileSpec("planning.xlsx#Teams"),
            "team_plan": FileSpec("planning.xlsx#Teamrooster"),
        },
        notes=(
            "The workbook sheets also have the title lines; the totals line is not a row.",
            "Column names keep the plant's English names; the layout is what is messy.",
        ),
    ),
)


# --- reading the plant's clean export -----------------------------------------------------------


def _mapping(name: str) -> dict[str, Any]:
    data: dict[str, Any] = yaml.safe_load((MAPPINGS / name / "mapping.yaml").read_text())
    return data


def _source_tables(case: Case) -> dict[str, tuple[list[str], list[dict[str, str]]]]:
    m = _mapping(case.mapping)
    folder = SYN / case.plant / "erp_export" / "clean" / "csv"
    tables = {}
    for path in sorted(folder.glob("*.csv")):
        with path.open(encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f, delimiter=m["delimiter"])
            rows = list(reader)
            tables[path.stem] = (list(reader.fieldnames or []), rows)
    return tables


def _column_rules(
    mapping: dict[str, Any],
) -> dict[str, dict[str, list[tuple[str, dict[str, Any]]]]]:
    """Source table → source column → [(target.field, rule)]."""
    out: dict[str, dict[str, list[tuple[str, dict[str, Any]]]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for t in mapping["tables"]:
        for fname, rule in t["fields"].items():
            if "from" in rule:
                out[t["source"]][rule["from"]].append((f"{t['target']}.{fname}", rule))
    return out


# --- writing a case -----------------------------------------------------------------------------


def _converter(
    case: Case, mapping: dict[str, Any], rules: list[tuple[str, dict[str, Any]]]
) -> tuple[str, Callable[[str], Any]]:
    kinds = {r.get("type", "str") for _, r in rules}
    tables = {r["code_table"] for _, r in rules if "code_table" in r}
    src_fmt = mapping["datetime_format"]
    src_dec = mapping.get("decimal", ".")

    if "datetime" in kinds:

        def dt(v: str) -> Any:
            return datetime.strptime(v, src_fmt) if v else None

        return "datetime", dt
    if kinds & {"decimal", "eur_cents"}:

        def num(v: str) -> Any:
            return float(v.replace(src_dec, ".")) if v else None

        return "number", num
    if "int" in kinds:
        return "int", lambda v: int(v) if v else None
    if tables:
        recodes = [case.recode[t] for t in tables if t in case.recode]

        def code(v: str) -> Any:
            if not v:
                return v
            if "list" in kinds:
                return ";".join(_recode(x, recodes) for x in v.split(";"))
            return _recode(v, recodes)

        return "code", code
    return "str", lambda v: v


def _recode(v: str, recodes: list[dict[str, str]]) -> str:
    for r in recodes:
        if v in r:
            return r[v]
    return v


def _text(case: Case, kind: str, value: Any) -> list[str]:
    if value is None or value == "":
        return [""] if not (kind == "datetime" and case.split_datetime) else ["", ""]
    if kind == "datetime":
        if case.split_datetime:
            return [value.strftime(case.split_datetime[0]), value.strftime(case.split_datetime[1])]
        return [value.strftime(case.datetime_format)]
    if kind == "number":
        return [f"{value:.2f}".replace(".", case.decimal)]
    return [str(value)]


def _header(case: Case, spec: FileSpec, col: str, kind: str) -> list[str]:
    if kind == "datetime" and case.split_datetime:
        names = case.split_names.get(col) or (f"{col} datum", f"{col} tijd")
        return list(names)
    return [spec.columns.get(col, col)]


def _csv_text(case: Case, header: list[str], rows: list[list[str]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=case.delimiter, lineterminator="\n")
    for line in case.title_lines:
        buf.write(line + "\n")
    w.writerow(header)
    w.writerows(rows)
    if case.footer:
        buf.write(f"Totaal{case.delimiter}{len(rows)} regels\n")
    return ("﻿" if case.bom else "") + buf.getvalue()


def _xlsx_bytes(case: Case, sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> bytes:
    wb = Workbook()
    for blank in list(wb.worksheets):
        wb.remove(blank)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        for line in case.title_lines:
            ws.append([line])
        ws.append(header)
        for r in rows:
            ws.append(r)
        if case.footer:
            ws.append(["Totaal", f"{len(rows)} regels"])
    wb.properties.created = XLSX_EPOCH
    wb.properties.modified = XLSX_EPOCH
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bookings(
    case: Case, cols: list[str], rows: list[dict[str, str]], mapping: dict[str, Any]
) -> tuple[list[str], list[list[str]]]:
    """Pair each start with the next finish of the same job, operation and step."""
    b = case.bookings or {}
    fmt = mapping["datetime_format"]
    started: dict[tuple[str, str, str], datetime] = {}
    out = []
    for r in rows:
        key = (r["Bon"], r["Bewerking"], r["Werkplek"])
        at = datetime.strptime(r["Tijdstip"], fmt)
        if r["Status"] == "GESTART":
            started[key] = at
        elif r["Status"] == "GEREED" and key in started:
            begin = started.pop(key)
            minutes = (at - begin) / timedelta(minutes=1)
            out.append(
                [
                    r["Bon"],
                    r["Bewerking"],
                    r["Werkplek"],
                    begin.strftime("%d-%m-%Y"),
                    begin.strftime("%H:%M"),
                    f"{minutes:.1f}".replace(".", case.decimal),
                    r["Aantal"],
                ]
            )
    header = [b[k] for k in ("job", "operation", "step", "date", "time", "duration", "quantity")]
    return header, out


def _flow(plant: str) -> dict[str, Any]:
    """Routes and arrows of the plant's jobs from its standard event log: each job's steps in
    order of their start touchpoints (a restart at the same step is one visit)."""
    with (SYN / plant / "events.csv").open(encoding="utf-8", newline="") as f:
        events = [r for r in csv.DictReader(f) if r["event_type"] == "start" and r["resource_code"]]
    starts: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for r in events:
        starts[r["order_id"]].append((r["at"], r["resource_code"]))
    routes: Counter[tuple[str, ...]] = Counter()
    arrows: Counter[tuple[str, str]] = Counter()
    for visits in starts.values():
        seq: list[str] = []
        for _, step in sorted(visits):
            if not seq or seq[-1] != step:
                seq.append(step)
        routes[tuple(seq)] += 1
        arrows.update(pairwise(seq))
    jobs = len(starts)
    leaving: Counter[str] = Counter()
    for (a, _), n in arrows.items():
        leaving[a] += n
    top = routes.most_common(TOP_ROUTES)
    return {
        "jobs": jobs,
        "rule": "steps of a job in order of their start touchpoints; a restart at the same step "
        "is one visit; rework is a step visited again later",
        "routes": [
            {
                "steps": list(r),
                "jobs": n,
                "share_pct": round(100 * n / jobs, 1),
                "rework": len(set(r)) < len(r),
            }
            for r, n in top
        ],
        "other_routes_jobs": jobs - sum(n for _, n in top),
        "arrows": [
            {"from": a, "to": b, "jobs": n, "share_of_leaving_pct": round(100 * n / leaving[a], 1)}
            for (a, b), n in sorted(arrows.items())
        ],
    }


def build_case(case: Case) -> tuple[dict[Path, bytes], dict[str, Any]]:
    mapping = _mapping(case.mapping)
    rules = _column_rules(mapping)
    tables = _source_tables(case)
    folder = OUT / case.name
    texts: dict[str, str] = {}
    sheets: dict[str, dict[str, tuple[list[str], list[list[Any]]]]] = defaultdict(dict)
    expected_files: dict[str, Any] = {}
    columns: dict[str, dict[str, Any]] = {}
    personal = set(mapping.get("personal_columns", []))

    for source, spec in case.files.items():
        cols, rows = tables[source]
        convs = {c: _converter(case, mapping, rules[source].get(c, [])) for c in cols}
        header = [h for c in cols for h in _header(case, spec, c, convs[c][0])]
        native = "#" in spec.name
        out_rows: list[list[Any]] = []
        for r in rows:
            line: list[Any] = []
            for c in cols:
                kind, conv = convs[c]
                value = conv(r[c])
                if native:
                    if kind == "datetime" and case.split_datetime:
                        line += [value.date() if value else None, value.time() if value else None]
                    elif isinstance(value, str) and value.isdigit() and not value.startswith("0"):
                        line.append(int(value))  # Excel holds order numbers as numbers
                    else:
                        line.append(None if value == "" else value)
                else:
                    line += _text(case, kind, value)
            out_rows.append(line)
        if native:
            book, sheet = spec.name.split("#")
            sheets[book][sheet] = (header, out_rows)
        else:
            texts[spec.name] = _csv_text(case, header, [[str(x) for x in r] for r in out_rows])

        targets = sorted({t["target"] for t in mapping["tables"] if t["source"] == source})
        expected_files[spec.name] = {"role": targets[0] if len(targets) == 1 else targets}
        cmap: dict[str, Any] = {}
        for c in cols:
            names = _header(case, spec, c, convs[c][0])
            if c in personal:
                match: Any = "personal (dropped on read, never shown or sent)"
            elif c in rules[source]:
                match = sorted(t for t, _ in rules[source][c])
            else:
                match = "not used"
            if isinstance(match, list) and len(match) == 1:
                match = match[0]
            if len(names) == 2:
                cmap[names[0]] = {"date part of": match}
                cmap[names[1]] = {"time part of": match}
            else:
                cmap[names[0]] = match
        columns[spec.name] = cmap

    if case.bookings:
        header, out_rows = _bookings(case, *tables["mes_events"], mapping)
        b = case.bookings
        texts[b["file"]] = _csv_text(case, header, out_rows)
        expected_files[b["file"]] = {"role": "hour_bookings"}
        columns[b["file"]] = {
            b["job"]: "events.order_id",
            b["operation"]: "events.operation_seq",
            b["step"]: "events.resource_code",
            b["date"]: {"date part of": "booking start"},
            b["time"]: {"time part of": "booking start"},
            b["duration"]: "booking duration_min",
            b["quantity"]: "events.quantity_units",
        }

    for name, extra in case.extra_files.items():
        texts[name] = extra
    for name, reason in case.not_used.items():
        expected_files[name] = {"role": "not_used", "reason": reason}

    files: dict[Path, bytes] = {folder / n: t.encode("utf-8") for n, t in texts.items()}
    for book, book_sheets in sheets.items():
        files[folder / book] = _xlsx_bytes(case, book_sheets)

    code_tables = {}
    for tname, table in mapping.get("code_tables", {}).items():
        used = any(
            r.get("code_table") == tname
            for t in mapping["tables"]
            if t["source"] in case.files
            for r in t["fields"].values()
        )
        if not used:
            continue
        rec = case.recode.get(tname, {})
        code_tables[tname] = {rec.get(k, k): v for k, v in table.items()}
    if case.mapping == "generic" and "resource" not in code_tables:
        code_tables["resource"] = mapping["code_tables"]["resource"]

    expected = {
        "case": case.name,
        "style": case.style,
        "rule": RULE,
        "plant": case.plant,
        "plant_answer": f"data/synthetic/{case.plant}/expected_answers.yaml",
        "answer": case.answer,
        "format": {
            "delimiter": {"\t": "tab"}.get(case.delimiter, case.delimiter),
            "decimal": case.decimal,
            "datetime": (
                {"date": case.split_datetime[0], "time": case.split_datetime[1]}
                if case.split_datetime
                else case.datetime_format
            ),
            "timezone": mapping["timezone"],
            "title_lines_above_header": len(case.title_lines),
            "totals_line_at_end": case.footer,
            "byte_order_mark": case.bom,
        },
        "files": dict(sorted(expected_files.items())),
        "columns": dict(sorted(columns.items())),
        "code_tables": code_tables,
        "flow": _flow(case.plant),
        "notes": list(case.notes),
    }
    files[folder / "expected.yaml"] = yaml.dump(
        expected, Dumper=_NoAliases, sort_keys=False, allow_unicode=True, width=100
    ).encode("utf-8")
    return files, expected


class _NoAliases(yaml.SafeDumper):
    """Write repeated values in full: expected.yaml is read by people."""

    def ignore_aliases(self, data: Any) -> bool:
        return True


def build() -> dict[Path, bytes]:
    """Every file of every case, by path. Pure: same input, same output."""
    files: dict[Path, bytes] = {}
    for case in CASES:
        files.update(build_case(case)[0])
    readme = (
        "# Mapping-wizard test exports (ADR-0011)\n\n"
        "Generated by `scripts/gen_wizard_cases.py` from the synthetic plants; "
        "never edit by hand.\n\n"
    )
    readme += "| Case | Plant | Style |\n|---|---|---|\n"
    readme += "".join(f"| `{c.name}` | {c.plant} | {c.style} |\n" for c in CASES)
    files[OUT / "README.md"] = readme.encode("utf-8")
    return files


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    built = build()
    stale = []
    if args.check:
        for path, content in built.items():
            if not path.exists() or not _same(path, content):
                stale.append(path.relative_to(ROOT))
        extra = {p for p in OUT.rglob("*") if p.is_file()} - set(built)
        stale += [p.relative_to(ROOT) for p in sorted(extra)]
    else:
        for path, content in built.items():
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    for path in stale:
        print(f"stale: {path}", file=sys.stderr)
    return 1 if stale else 0


def _same(path: Path, content: bytes) -> bool:
    if path.suffix == ".xlsx":
        return _sheets(path.read_bytes()) == _sheets(content)
    return path.read_bytes() == content


def _sheets(data: bytes) -> list[list[tuple[Any, ...]]]:
    """Cell values of every sheet: the zip itself carries timestamps."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True)
    return [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]


if __name__ == "__main__":
    sys.exit(main())
