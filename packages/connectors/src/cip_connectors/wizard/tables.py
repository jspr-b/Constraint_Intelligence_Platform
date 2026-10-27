"""Read a whole export as it comes (ADR-0011 §2): CSV, TXT or TSV files and every sheet of an
Excel workbook, with title lines above the header and a totals line below the data.

Personal columns are dropped as each table is read; their values never enter a row. A file or
sheet that cannot be read says why with a stable ``problem`` code, and the rest of the export
continues.
"""

import csv
import io
import re
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from cip_connectors.convert import Cell
from cip_connectors.privacy import personal_column_rule
from cip_connectors.readers import SEPARATORS, DropRule, UnreadableFileError, decode

TEXT_SUFFIXES = (".csv", ".txt", ".tsv")
EXCEL_SUFFIXES = (".xlsx",)
SCAN_LINES = 40
"""Lines read to find the separator and the header."""
TOTAL_WORDS = re.compile(
    r"^\s*(totaal|total|totals|summe|gesamt|som|sum|count|aantal regels)\b", re.I
)


@dataclass
class ExportTable:
    """One table of the export: a file, or one sheet of a workbook."""

    id: str
    """``file`` or ``file#sheet``: how the client and the stored mapping name this table."""
    file: str
    sheet: str | None
    header: list[str] = field(default_factory=list)
    """Column names, personal columns left out."""
    rows: list[dict[str, Cell]] = field(default_factory=list)
    personal_columns: list[str] = field(default_factory=list)
    title_lines: int = 0
    """Lines above the header (a report title, an export date)."""
    totals_lines: int = 0
    """Lines at the end that are totals, not rows."""
    encoding: str | None = None
    byte_order_mark: bool = False
    """The file starts with a UTF-8 byte-order mark (Excel's "CSV UTF-8")."""
    delimiter: str | None = None
    problem: str | None = None
    """empty, no_header, not_text, single_column, unreadable_excel, not_a_table_file."""

    @property
    def readable(self) -> bool:
        return self.problem is None


def read_export(paths: list[Path]) -> list[ExportTable]:
    """Every table of the export, in file order; sheets in workbook order."""
    drop = personal_column_rule()
    tables: list[ExportTable] = []
    for path in sorted(paths, key=lambda p: p.name.lower()):
        suffix = path.suffix.lower()
        if suffix in EXCEL_SUFFIXES:
            tables += _read_workbook(path, drop)
        elif suffix in TEXT_SUFFIXES:
            tables.append(_read_text(path, drop))
        else:
            tables.append(ExportTable(path.name, path.name, None, problem="not_a_table_file"))
    return tables


def _fields(line: str, delimiter: str) -> list[str]:
    return next(csv.reader([line], delimiter=delimiter), [])


def _choose_delimiter(lines: list[str]) -> tuple[str, int]:
    """The separator that splits most lines into the same number of fields (> 1), and that
    number. Title lines above the header have one field, so they do not count against it."""
    best, best_count, best_score = SEPARATORS[0], 1, 0
    for d in SEPARATORS:
        counts = Counter(len(_fields(line, d)) for line in lines if line.strip())
        width, seen = max(counts.items(), key=lambda kv: (kv[1], kv[0]), default=(1, 0))
        if width > 1 and seen > best_score:
            best, best_count, best_score = d, width, seen
    return best, best_count


def _is_header(values: list[str]) -> bool:
    """A header row: two or more names, none of them a number or a date."""
    filled = [v.strip() for v in values if v.strip()]
    return len(filled) >= 2 and not any(re.fullmatch(r"[\d.,:/ -]+", v) for v in filled)


def _read_text(path: Path, drop: DropRule) -> ExportTable:
    table = ExportTable(path.name, path.name, None)
    raw = path.read_bytes()
    table.byte_order_mark = raw.startswith(b"\xef\xbb\xbf")
    try:
        text, table.encoding = decode(raw)
    except UnreadableFileError as exc:
        table.problem = exc.problem
        return table
    lines = text.lstrip("﻿").splitlines()
    if not any(line.strip() for line in lines):
        table.problem = "empty"
        return table
    head = lines[:SCAN_LINES]
    table.delimiter, width = _choose_delimiter(head)
    if width < 2:
        table.problem = "single_column"
        return table
    start = next(
        (
            i
            for i, line in enumerate(head)
            if len(_fields(line, table.delimiter)) == width
            and _is_header(_fields(line, table.delimiter))
        ),
        None,
    )
    if start is None:
        table.problem = "no_header"
        return table
    table.title_lines = start
    reader = csv.reader(io.StringIO("\n".join(lines[start:])), delimiter=table.delimiter)
    header = [h.strip() for h in next(reader)]
    data = [values for values in reader if any(v.strip() for v in values)]
    _fill(table, header, data, drop)
    return table


def _read_workbook(path: Path, drop: DropRule) -> list[ExportTable]:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:  # corrupt, password-protected or not an Excel file
        return [ExportTable(path.name, path.name, None, problem="unreadable_excel")]
    tables = []
    for ws in wb.worksheets:
        table = ExportTable(f"{path.name}#{ws.title}", path.name, ws.title, encoding="xlsx")
        for merged in list(ws.merged_cells.ranges):  # every cell gets the top-left value
            value = ws.cell(merged.min_row, merged.min_col).value
            ws.unmerge_cells(str(merged))
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    ws.cell(r, c).value = value
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        start = next(
            (i for i, r in enumerate(rows[:SCAN_LINES]) if _is_header([_str(v) for v in r])), None
        )
        if not any(any(v not in (None, "") for v in r) for r in rows):
            table.problem = "empty"
        elif start is None:
            table.problem = "no_header"
        else:
            table.title_lines = start
            header = [_str(v).strip() for v in rows[start]]
            data = [r for r in rows[start + 1 :] if any(v not in (None, "") for v in r)]
            _fill(table, header, data, drop)
        tables.append(table)
    return tables


def _fill(
    table: ExportTable, header: list[str], data: Sequence[Sequence[object]], drop: DropRule
) -> None:
    """Rows as column → value, without personal columns and without a totals line at the end."""
    while data and _is_totals(data[-1], len(header)):
        data = data[:-1]
        table.totals_lines += 1
    keep = [i for i, name in enumerate(header) if name and not drop(name)]
    table.personal_columns = [n for i, n in enumerate(header) if n and i not in keep]
    table.header = [header[i] for i in keep]
    table.rows = [
        {header[i]: _cell(values[i]) if i < len(values) else None for i in keep} for values in data
    ]


def _is_totals(values: Sequence[object], width: int) -> bool:
    filled = [v for v in values if v not in (None, "")]
    first = _str(filled[0]) if filled else ""
    return bool(TOTAL_WORDS.match(first)) and len(filled) <= max(2, width // 3)


def _str(value: object) -> str:
    return "" if value is None else str(value)


def _cell(value: object) -> Cell:
    if value is None or isinstance(value, str | int | float | datetime):
        return value.strip() if isinstance(value, str) else value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value)
