"""Mapping-wizard test exports (ADR-0011, CI0005 sequence 01).

The first tests check that every case and its `expected.yaml` agree with each other. Then the
wizard's known answers: file roles and links (sequence 02). The `xfail(strict=True)` tests are
the answers still to build; pytest reports them as passing unexpectedly once they work, so the
marker is removed in the sequence that turns them green.
"""

import csv
import io
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from cip_connectors.wizard import recognise_export

ROOT = Path(__file__).resolve().parents[3]
CASES = ROOT / "data" / "synthetic" / "wizard_cases"
NAMES = sorted(p.name for p in CASES.iterdir() if (p / "expected.yaml").exists())
PERSONAL = "personal (dropped on read, never shown or sent)"


def _expected(name: str) -> dict[str, Any]:
    data: dict[str, Any] = yaml.safe_load((CASES / name / "expected.yaml").read_text())
    return data


def _headers(name: str, file: str, fmt: dict[str, Any]) -> list[str]:
    """Header of a CSV file or a workbook sheet (``book.xlsx#sheet``), below any title lines."""
    skip = fmt["title_lines_above_header"]
    if "#" in file:
        book, sheet = file.split("#")
        ws = load_workbook(CASES / name / book, read_only=True)[sheet]
        rows = list(ws.iter_rows(min_row=skip + 1, max_row=skip + 1, values_only=True))
        return [str(c) for c in rows[0] if c is not None]
    text = (CASES / name / file).read_text(encoding="utf-8-sig")
    delimiter = {"tab": "\t"}.get(fmt["delimiter"], fmt["delimiter"])
    lines = text.splitlines()[skip:]
    return next(csv.reader(io.StringIO(lines[0]), delimiter=delimiter))


def test_there_are_six_cases_from_four_plants() -> None:
    assert NAMES == ["exact_en", "isah_xlsx", "mes_us", "messy", "ridder_hours", "sap_de"]
    assert {_expected(n)["plant"] for n in NAMES} == {"plant_a", "plant_b", "plant_c", "plant_d"}


@pytest.mark.parametrize("name", NAMES)
def test_every_expected_column_is_in_its_file(name: str) -> None:
    exp = _expected(name)
    for file, columns in exp["columns"].items():
        assert set(columns) <= set(_headers(name, file, exp["format"])), (name, file)


@pytest.mark.parametrize("name", NAMES)
def test_every_case_has_a_known_answer_and_its_needed_files(name: str) -> None:
    exp = _expected(name)
    assert (ROOT / exp["plant_answer"]).exists()
    roles = {
        r
        for f in exp["files"].values()
        for r in ([f["role"]] if isinstance(f["role"], str) else f["role"])
    }
    assert roles & {"events", "hour_bookings"}, "an event log or hour bookings"
    assert {"resources", "calendar"} <= roles, "a steps list and a shift sheet"
    assert exp["flow"]["routes"][0]["share_pct"] > 20, "a main route the canvas can show"


@pytest.mark.parametrize("name", [n for n in NAMES if n != "ridder_hours"])
def test_the_personal_column_is_in_the_file_and_marked(name: str) -> None:
    exp = _expected(name)
    personal = [c for cols in exp["columns"].values() for c, m in cols.items() if m == PERSONAL]
    assert len(personal) == 1, "each event log carries one personal column, to be dropped"


def test_hour_bookings_have_a_lower_confidence_answer() -> None:
    exp = _expected("ridder_hours")
    assert exp["files"]["Urenregistratie.csv"]["role"] == "hour_bookings"
    assert exp["answer"]["constraint"] == "PNT-01"
    assert "lower than high" in exp["answer"]["confidence"]


# --- the wizard's known answers ----------------------------------------------------------------


def _files(name: str) -> list[Path]:
    return sorted(p for p in (CASES / name).iterdir() if p.name != "expected.yaml")


@pytest.mark.parametrize("name", NAMES)
def test_wizard_sorts_the_files(name: str) -> None:
    found = recognise_export(_files(name)).summary()["files"]
    expected = {
        f: {"role": None if v["role"] == "not_used" else v["role"]}
        for f, v in _expected(name)["files"].items()
    }
    assert found == expected


@pytest.mark.parametrize("name", NAMES)
def test_wizard_sorts_the_files_by_content_alone(name: str, tmp_path: Path) -> None:
    """Files renamed table0, table1, ...: the roles hold, only less sure."""
    renamed = {}
    for i, f in enumerate(_files(name)):
        target = tmp_path / f"table{i}{f.suffix}"
        target.write_bytes(f.read_bytes())
        renamed[target.name] = f.name
    found = recognise_export(sorted(tmp_path.iterdir()))
    expected = _expected(name)["files"]
    files = found.summary()["files"]
    for r in found.roles:
        file, _, sheet = r.table.partition("#")
        original = renamed[file] + (f"#{sheet}" if sheet else "")
        want = expected[original]["role"]
        if want == "not_used":
            assert r.role is None, original
        else:
            # A sheet keeps its name inside the renamed workbook, so it may still be "sure".
            sure = ("likely", "sure") if sheet else ("likely",)
            assert files[r.table]["role"] == want and r.sure in sure, original


@pytest.mark.parametrize("name", NAMES)
def test_wizard_links_the_event_log_to_the_steps_list(name: str) -> None:
    exp = _expected(name)
    log = next(f for f, v in exp["files"].items() if v["role"] in ("events", "hour_bookings"))
    steps = next(f for f, v in exp["files"].items() if v["role"] == ["resources", "stations"])

    def column(file: str, field: str) -> str:
        return str(
            next(
                c
                for c, m in exp["columns"][file].items()
                if m == field or (isinstance(m, list) and field in m)
            )
        )

    want = (
        log,
        column(log, "events.resource_code"),
        steps,
        column(steps, "resources.resource_code"),
    )
    links = recognise_export(_files(name)).links
    assert want in {(k.table, k.column, k.to_table, k.to_column) for k in links}


@pytest.mark.parametrize("name", NAMES)
def test_wizard_matches_every_column_by_rules(name: str) -> None:
    """Sequence 03: every column of every placed table gets the field of the plant's true
    mapping from rules alone, and personal columns are only listed as dropped."""
    exp = _expected(name)
    summary = recognise_export(_files(name)).summary()
    placed = {t: c for t, c in summary["columns"].items() if summary["files"][t]["role"]}
    assert placed == exp["columns"]


ALIAS = {
    "event_status": "events.event_type",
    "action": "events.event_type",
    "order_status": "orders.status",
    "job_status": "orders.status",
    "calendar_kind": "calendar.kind",
    "window": "calendar.kind",
    "step_kind": "resources.kind",
    "resource": "resources.resource_code",
    "station": "stations.station_code",
}
"""The code tables of the two hand-written mappings, by the field they fill."""


@pytest.mark.parametrize("name", NAMES)
def test_wizard_gives_every_code_in_the_data_its_meaning(name: str) -> None:
    """Sequence 04: each status code in the export means what the plant's true mapping says,
    and no code is left without a meaning. A known answer may list codes this export does not
    use (SAP's ABG); only the codes in the data count."""
    found = recognise_export(_files(name))
    tables = found.code_tables()
    for table, want in _expected(name)["code_tables"].items():
        f = ALIAS[table]
        if f in ("resources.resource_code", "stations.station_code"):
            # Steps and stations keep the plant's own codes (ADR-0011 §4): the same codes.
            assert set(tables[f]) == set(want), (name, table)
            assert all(code == own for code, own in tables[f].items())
            continue
        in_data = {c.code for c in found.codes if c.target == f}
        assert in_data, (name, table)
        assert in_data <= set(want), (name, table, in_data - set(want))
        assert tables.get(f, {}) == {k: v for k, v in want.items() if k in in_data}, (name, table)


@pytest.mark.parametrize("name", NAMES)
def test_wizard_reads_the_export_format(name: str) -> None:
    """Separator, decimal mark, date-time format, title and totals lines and the byte-order
    mark, from the files alone; the time zone is the plant's. A workbook has no separator."""
    exp = _expected(name)["format"]
    found = recognise_export(_files(name)).summary(timezone=exp["timezone"])["format"]
    if all(f.endswith(".xlsx") for f in _expected(name)["files"]) or all(
        "#" in f for f in _expected(name)["files"]
    ):
        assert found.pop("delimiter") is None
        exp = {k: v for k, v in exp.items() if k != "delimiter"}
    assert found == exp
