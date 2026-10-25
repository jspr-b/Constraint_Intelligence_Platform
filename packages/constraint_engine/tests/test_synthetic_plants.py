"""Lock the committed synthetic plants so pillar tests can trust them."""

from pathlib import Path

import pytest
import yaml

from constraint_engine.io.plant_loader import load_plant
from constraint_engine.synthetic.config import load_config
from constraint_engine.synthetic.expected import ExpectedAnswers
from constraint_engine.synthetic.model import simulate
from constraint_engine.synthetic.writer import write_plant

DATA = Path(__file__).resolve().parents[3] / "data" / "synthetic"
PLANTS = ["plant_a", "plant_b", "plant_c", "plant_d"]


def answers(plant: str) -> ExpectedAnswers:
    return ExpectedAnswers.model_validate(
        yaml.safe_load((DATA / plant / "expected_answers.yaml").read_text())
    )


@pytest.mark.slow
@pytest.mark.parametrize("plant", PLANTS)
def test_regenerating_reproduces_committed_files(plant: str, tmp_path: Path) -> None:
    out = write_plant(simulate(load_config(DATA / plant / "config.yaml")), tmp_path / plant)
    committed = {
        p.name
        for p in (DATA / plant).iterdir()
        if p.is_file() and p.name not in ("config.yaml", "README.md", "known_answer.yaml")
    }
    assert {p.name for p in out.iterdir() if p.is_file()} == committed
    for name in committed:
        assert (out / name).read_bytes() == (DATA / plant / name).read_bytes(), name


@pytest.mark.parametrize("plant", PLANTS)
def test_plant_loads_without_errors(plant: str) -> None:
    loaded = load_plant(DATA / plant)
    assert loaded.row_errors == []
    assert loaded.missing_files == {}
    assert len(loaded.events) > 1000


@pytest.mark.parametrize("plant", PLANTS)
def test_expected_answers_validate(plant: str) -> None:
    assert answers(plant).plant_code == f"PLANT-{plant[-1].upper()}"


def test_plant_a_constraint_is_paint() -> None:
    a = answers("plant_a")
    assert a.constraint == "PNT-01" and a.constraint_kind == "internal"
    assert a.throughput_eur_cents_per_constraint_h["R-8"] is None  # rails use no paint time
    assert all(
        month.breaks_h > 0 and month.setup_h > 0 for month in a.lost_constraint_h_by_month.values()
    )


def test_plant_b_is_market_bound() -> None:
    b = answers("plant_b")
    assert b.constraint == "none" and b.constraint_kind == "market"
    assert max(b.utilisation_pct.values()) < 85


def test_plant_c_constraint_moves_to_wld02_in_orderbook_week_7() -> None:
    c = answers("plant_c")
    assert c.constraint == "PNT-01"
    assert c.constraint_moves is not None
    assert c.constraint_moves["to"] == "WLD-02" and c.constraint_moves["orderbook_week"] == 7


ERP = DATA / "plant_a" / "erp_export"
ERP_TABLES = {
    "werkplekken",
    "artikelen",
    "bewerkingen",
    "orders",
    "mes_events",
    "rooster",
    "overuren",
    "capaciteitskosten",
    "omsteltijden",
    "uitbesteding",
}


def xlsx_values(path: Path) -> list[tuple[object, ...]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True).active
    assert ws is not None
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


@pytest.mark.slow
def test_erp_export_regenerates(tmp_path: Path) -> None:
    from constraint_engine.synthetic.erp_export import export

    out = export(DATA / "plant_a", tmp_path / "erp")
    for variant in ("clean", "defects"):
        for name in ERP_TABLES:
            csv_path = Path(variant) / "csv" / f"{name}.csv"
            assert (out / csv_path).read_bytes() == (ERP / csv_path).read_bytes(), csv_path
            xlsx_path = Path(variant) / "xlsx" / f"{name}.xlsx"
            assert xlsx_values(out / xlsx_path) == xlsx_values(ERP / xlsx_path), xlsx_path
    assert (out / "defects" / "defects.json").read_text() == (
        ERP / "defects" / "defects.json"
    ).read_text()


def test_erp_export_has_every_table_in_csv_and_excel() -> None:
    for variant in ("clean", "defects"):
        assert {p.stem for p in (ERP / variant / "csv").glob("*.csv")} == ERP_TABLES
        assert {p.stem for p in (ERP / variant / "xlsx").glob("*.xlsx")} == ERP_TABLES


def test_employee_column_only_holds_fake_test_values() -> None:
    header, *rows = (ERP / "clean" / "csv" / "mes_events.csv").read_text().splitlines()
    col = header.split(";").index("Medewerker")
    assert {r.split(";")[col] for r in rows} <= {f"TEST-MEDEWERKER-{n:02d}" for n in range(1, 7)}


def test_defect_counts_match_files() -> None:
    import json

    defects = json.loads((ERP / "defects" / "defects.json").read_text())
    clean = (ERP / "clean" / "csv" / "mes_events.csv").read_text().splitlines()
    broken = (ERP / "defects" / "csv" / "mes_events.csv").read_text().splitlines()
    assert len(clean) - len(broken) == defects["missing_end_events"] + defects["scrap_rows_removed"]
    assert (
        sum(";LAS9;" in f";{line}" or line.startswith("LAS9;") for line in broken)
        == defects["unknown_machine_code_rows"]
    )


def test_plant_d_meets_its_answer_fixed_before_any_code() -> None:
    """known_answer.yaml was written on 2026-10-23 before the generator could make plant D."""
    fixed = yaml.safe_load((DATA / "plant_d" / "known_answer.yaml").read_text())
    d = answers("plant_d")
    assert d.constraint == fixed["constraint"] and d.constraint_kind == fixed["constraint_kind"]
    assert d.constraint_type == fixed["constraint_type"]
    assert d.constraint_covers == fixed["constraint_covers"]
    assert d.forgotten_end_scan_share_pct is not None
    assert abs(d.forgotten_end_scan_share_pct - fixed["forgotten_finish_scan_share_pct"]) < 1.0
    top = d.utilisation_pct[d.constraint]
    for code in fixed["not_constraint"]:
        assert d.utilisation_pct[code] < top - 20, code
    assert fixed["money_source"] == "job_value" and d.throughput_eur_cents_per_constraint_h["jobs"]
    assert not (DATA / "plant_d" / "products.csv").exists()
    assert not (DATA / "plant_d" / "routings.csv").exists()


def test_plant_d_shares_no_step_numbers_and_one_team() -> None:
    d = load_plant(DATA / "plant_d")
    assert all(e.operation_seq is None for e in d.events)
    assert all(o.sku is None and o.order_value_eur_cents for o in d.orders)
    assert [(t.team_code, t.resource_codes) for t in d.teams] == [
        ("FINISH-TEAM", ("FINISH", "PACK"))
    ]
    assert {r.kind for r in d.resources} == {"machine", "manual"}
